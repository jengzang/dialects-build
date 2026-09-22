import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.yindian_llm import (
    PROMPT,
    PROMPT_VERSION,
    _input_fingerprint,
    _load_api_key,
    run_analysis,
)


class YindianLlmTests(unittest.TestCase):
    def _source(self, root: Path) -> Path:
        source_dir = root / "pull_yindian"
        source_dir.mkdir()
        (source_dir / "甲地.tsv").write_text(
            "#漢字\t音標\t解釋\n"
            "□\tka1\t~腰:脊背微駝\n"
            "一\ti1\t\n"
            "二\tki2\t又,門~:門栓\n"
            "三\tsa3\t单纯读音说明\n",
            encoding="utf-8",
        )
        return source_dir

    def test_loads_deepseek_key_from_dotenv_with_hyphenated_name(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            dotenv_path = Path(tmpdir) / ".env"
            dotenv_path.write_text(
                'deepseek-api-key="test-secret"\n',
                encoding="utf-8",
            )

            self.assertEqual(
                _load_api_key(dotenv_path=dotenv_path),
                "test-secret",
            )

    def test_prompt_requires_expanding_tilde_into_the_original_character(self):
        self.assertIn("`~` 或 `～` 是音典注释中的占位符", PROMPT)
        self.assertIn("䂿背脊", PROMPT)
        self.assertIn("不能只输出 `䂿`", PROMPT)
        self.assertIn('"置信度":0.9', PROMPT)
        self.assertNotIn("不要把解释文字硬拼成方言词", PROMPT)

    def test_prompt_version_is_part_of_cached_result_fingerprint(self):
        record = {
            "id": "南寧.tsv:195",
            "原始汉字": "䂿",
            "原始注释": "(白)砸、捶,~背脊,俗字撘",
            "来源文件": "南寧.tsv",
            "原始行号": 195,
        }

        current_fingerprint = _input_fingerprint(record)
        original_version = PROMPT_VERSION
        try:
            import scripts.yindian_llm as yindian_llm

            yindian_llm.PROMPT_VERSION = "different-prompt"
            self.assertNotEqual(current_fingerprint, _input_fingerprint(record))
        finally:
            yindian_llm.PROMPT_VERSION = original_version

    def test_sends_each_annotated_row_without_audio_or_mechanical_candidates(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            seen = []

            def fake_request(point, batch):
                seen.extend(batch)
                return {
                    "results": [
                        {
                            "id": item["id"],
                            "entries": [],
                        }
                        for item in batch
                    ]
                }

            run_analysis(
                self._source(root),
                points=["甲地"],
                batch_size=2,
                workers=1,
                output_dir=root / "output",
                request_fn=fake_request,
            )

            self.assertEqual([item["id"] for item in seen], [
                "甲地.tsv:2",
                "甲地.tsv:4",
                "甲地.tsv:5",
            ])
            self.assertEqual(
                set(seen[0]),
                {"id", "原始汉字", "原始注释", "来源文件", "原始行号"},
            )
            self.assertNotIn("音标", seen[0])
            self.assertNotIn("规则生成候选", seen[0])

    def test_batches_are_independent_and_writes_final_xlsx(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            calls = []

            def fake_request(point, batch):
                calls.append(batch)
                return {
                    "results": [
                        {
                            "id": item["id"],
                            "entries": [
                                {
                                    "方言词": "方" + item["原始汉字"],
                                    "标准词": ["标准一", "标准二"]
                                    if item["原始汉字"] == "□"
                                    else [],
                                    "置信度": 0.8,
                                }
                            ]
                            if item["原始汉字"] != "三"
                            else [],
                        }
                        for item in batch
                    ]
                }

            result = run_analysis(
                self._source(root),
                points=["甲地"],
                batch_size=2,
                workers=1,
                output_dir=root / "output",
                request_fn=fake_request,
            )

            self.assertEqual(result["甲地"]["rows"], 3)
            self.assertEqual(result["甲地"]["batches"], 2)
            self.assertEqual([len(batch) for batch in calls], [2, 1])

            xlsx_path = root / "output" / "xlsx" / "甲地.xlsx"
            self.assertTrue(xlsx_path.exists())
            workbook = load_workbook(xlsx_path)
            self.assertEqual(workbook.sheetnames, ["词表", "逐行结果"])

            sheet = workbook["词表"]
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(
                headers,
                [
                    "标准词",
                    "方言词",
                    "音标",
                    "释义",
                    "类别",
                    "置信度",
                    "原始汉字",
                    "原字音标",
                    "原始注释",
                    "来源文件",
                    "原始行号",
                ],
            )
            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(rows), 3)
            self.assertEqual(rows[0][0:6], ("标准一", "方□", None, "~腰:脊背微駝", None, 0.8))
            self.assertEqual(rows[1][0:2], ("标准二", "方□"))
            self.assertEqual(rows[2][0:2], (None, "方二"))
            self.assertEqual(rows[0][7], "ka1")

            row_sheet = workbook["逐行结果"]
            row_headers = [cell.value for cell in row_sheet[1]]
            self.assertEqual(
                row_headers,
                [
                    "id",
                    "原始汉字",
                    "原始注释",
                    "来源文件",
                    "原始行号",
                    "状态",
                    "LLM结果",
                ],
            )
            self.assertEqual(row_sheet.max_row, 4)

            result_lines = (
                root / "output" / "result" / "甲地.jsonl"
            ).read_text(encoding="utf-8").splitlines()
            self.assertEqual(len(result_lines), 3)
            first_result = json.loads(result_lines[0])
            self.assertEqual(first_result["id"], "甲地.tsv:2")
            self.assertRegex(first_result["input_fingerprint"], r"^[0-9a-f]{64}$")

    def test_persists_completed_batches_when_a_later_batch_fails(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            calls = 0

            def flaky_request(point, batch):
                nonlocal calls
                calls += 1
                if calls == 2:
                    raise RuntimeError("temporary failure")
                return {
                    "results": [
                        {"id": item["id"], "entries": []}
                        for item in batch
                    ]
                }

            summary = run_analysis(
                self._source(root),
                points=["甲地"],
                batch_size=2,
                workers=1,
                output_dir=root / "output",
                request_fn=flaky_request,
            )

            self.assertEqual(summary["甲地"]["status"], "partial")
            self.assertEqual(summary["甲地"]["completed"], 2)
            self.assertEqual(summary["甲地"]["errors"], 1)
            result_lines = (
                root / "output" / "result" / "甲地.jsonl"
            ).read_text(encoding="utf-8").splitlines()
            self.assertEqual(len(result_lines), 2)
            error_lines = (
                root / "output" / "result" / "甲地.errors.jsonl"
            ).read_text(encoding="utf-8").splitlines()
            self.assertEqual(len(error_lines), 1)
            manifest = json.loads(
                (root / "output" / "manifest.json").read_text(encoding="utf-8")
            )
            self.assertEqual(manifest["甲地"]["status"], "partial")

    def test_changed_source_row_does_not_reuse_cached_result(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source_dir = root / "pull_yindian"
            source_dir.mkdir()
            source_path = source_dir / "甲地.tsv"
            source_path.write_text(
                "#漢字\t音標\t解釋\n□\tka1\t第一版注释\n",
                encoding="utf-8",
            )
            calls = []

            def fake_request(point, batch):
                calls.append(batch)
                return {
                    "results": [
                        {"id": item["id"], "entries": []}
                        for item in batch
                    ]
                }

            run_analysis(
                source_dir,
                points=["甲地"],
                batch_size=1,
                workers=1,
                output_dir=root / "output",
                request_fn=fake_request,
            )
            source_path.write_text(
                "#漢字\t音標\t解釋\n□\tka1\t第二版注释\n",
                encoding="utf-8",
            )
            run_analysis(
                source_dir,
                points=["甲地"],
                batch_size=1,
                workers=1,
                output_dir=root / "output",
                request_fn=fake_request,
            )

            self.assertEqual(len(calls), 2)
            self.assertEqual(calls[1][0]["原始注释"], "第二版注释")

    def test_successful_retry_removes_stale_error_file(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            attempts = 0

            def flaky_once(point, batch):
                nonlocal attempts
                attempts += 1
                if attempts == 2:
                    raise RuntimeError("temporary failure")
                return {
                    "results": [
                        {"id": item["id"], "entries": []}
                        for item in batch
                    ]
                }

            output_dir = root / "output"
            source_dir = self._source(root)
            first_summary = run_analysis(
                source_dir,
                points=["甲地"],
                batch_size=2,
                workers=1,
                output_dir=output_dir,
                request_fn=flaky_once,
            )
            self.assertEqual(first_summary["甲地"]["status"], "partial")

            def successful_request(point, batch):
                return {
                    "results": [
                        {"id": item["id"], "entries": []}
                        for item in batch
                    ]
                }

            second_summary = run_analysis(
                source_dir,
                points=["甲地"],
                batch_size=2,
                workers=1,
                output_dir=output_dir,
                request_fn=successful_request,
            )
            self.assertEqual(second_summary["甲地"]["status"], "ok")
            self.assertFalse(
                (output_dir / "result" / "甲地.errors.jsonl").exists()
            )


if __name__ == "__main__":
    unittest.main()
