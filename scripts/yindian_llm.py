"""Use DeepSeek to turn annotated character-table rows into per-location XLSX word lists."""

import argparse
import concurrent.futures
import csv
import hashlib
import json
import os
import random
import re
import shutil
import tempfile
import threading
import time
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any, Callable, Optional

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE_DIR = PROJECT_ROOT / "data" / "raw" / "pull_yindian"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "data" / "raw" / "extract" / "word_llm"
DEFAULT_POINTS = ("泉州", "蘇州東村", "南寧")
API_URL = "https://api.deepseek.com/chat/completions"
PROMPT_VERSION = "2026-09-22-tilde-reconstruction-v2"

HANZI_HEADERS = frozenset({"漢字", "汉字"})
NOTE_HEADERS = frozenset({"解釋", "解释", "注釋", "注释", "notes", "note"})
TONE_HEADERS = frozenset({"音標", "音标"})

PROMPT = """你是汉语方言词汇整理者。请逐条独立分析输入数组中的记录。
每一条只能根据自己的“原始汉字”和“原始注释”判断，不能利用同一批次其他记录推断。

任务：
1. 判断这一行能否抽出一个或多个方言词；
2. 为每个方言词提供现代标准汉语词候选；
3. 给每个方言词一个 0 到 1 的置信度。

重要规则：
1. `~` 或 `～` 是音典注释中的占位符，表示本行的“原始汉字”，不是省略号，也不能忽略。
2. 每出现一个占位符，都要把它替换成原始汉字，并保留它所在的完整词语或短语。注释中的 `~背脊` 且原始汉字为 `䂿` 时，方言词必须写成 `䂿背脊`，不能只输出 `䂿`。
3. 真实案例：原始汉字为 `䂿`、原始注释为 `(白)砸、捶,~背脊,俗字撘` 时，至少返回
   `{"方言词":"䂿背脊","标准词":["砸背脊","捶背脊"],"置信度":0.9}`。
   如果同一行还明确说明原始汉字单独成词，可以另外返回单字词，但不能遗漏 `䂿背脊`。
4. 括号中的“白”“文”“又”等是读法或用法标签，不是词语本身；“俗字、本字、字形”说明不是词条，除非同一行明确给出完整词语。
5. 逗号、分号等可能分隔多个义项或词语。不要因为一个原始汉字可以解释为某个标准词，就丢掉注释中由占位符明确写出的更长方言词。
6. 纯读音、文白读、又音或音韵说明，如果不能形成明确词语，返回空 entries。
7. 不要猜测音标，不要输出音标；也不要凭空把注释中没有表达的文字拼成词。但按上述占位符规则还原出来的完整词语不是猜测，必须保留。
8. 一行可以有多个方言词；一个方言词可以有多个标准词候选。标准词应与该完整方言词的意义对应。

只返回 JSON 对象：
{"results":[
  {"id":"输入id","entries":[
    {"方言词":"...","标准词":["..."],"置信度":0.0}
  ]}
]}
没有明确词语时使用 "entries":[]。
"""

WORD_SHEET_HEADERS = [
    "标准词",
    "方言词",
    "音标",
    "释义",
    "类别",
    "置信度",
    "原始汉字",
    "原字音标",
    "原始注释",
    "来源文件",
    "原始行号",
]
ROW_SHEET_HEADERS = [
    "id",
    "原始汉字",
    "原始注释",
    "来源文件",
    "原始行号",
    "状态",
    "LLM结果",
]


def _load_api_key(
    api_key: Optional[str] = None,
    *,
    dotenv_path: Path = PROJECT_ROOT / ".env",
) -> Optional[str]:
    if api_key:
        return api_key

    environment_key = os.environ.get("DEEPSEEK_API_KEY")
    if environment_key:
        return environment_key

    if not dotenv_path.exists():
        return None

    with dotenv_path.open("r", encoding="utf-8") as dotenv_file:
        for line in dotenv_file:
            stripped = line.strip()
            if not stripped or stripped.startswith("#") or "=" not in stripped:
                continue
            name, value = stripped.split("=", 1)
            normalized_name = name.strip().removeprefix("export ").upper()
            normalized_name = normalized_name.replace("-", "_")
            if normalized_name != "DEEPSEEK_API_KEY":
                continue
            value = value.strip()
            if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
                value = value[1:-1]
            return value or None
    return None


def _normalize_header(value: str) -> str:
    return value.lstrip("\ufeff#").strip().casefold()


def _find_column(header: list[str], names: frozenset[str]) -> Optional[int]:
    wanted = {_normalize_header(name) for name in names}
    for index, value in enumerate(header):
        if _normalize_header(value) in wanted:
            return index
    return None


def _cell(row: list[str], index: Optional[int]) -> str:
    if index is None or index >= len(row):
        return ""
    return row[index]


def _is_character_table(header: list[str], hanzi_index: Optional[int]) -> bool:
    if hanzi_index is None:
        return False
    return header[hanzi_index].lstrip("\ufeff").strip() in {"#漢字", "#汉字"}


def _clear_dir(path: Path) -> None:
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True, exist_ok=True)


def _source_records(source_dir: Path, point: str) -> list[dict[str, Any]]:
    source_path = source_dir / f"{point}.tsv"
    if not source_path.exists():
        raise FileNotFoundError(source_path)

    records = []
    with source_path.open("r", encoding="utf-8-sig", newline="") as source_file:
        reader = csv.reader(source_file, delimiter="\t")
        header = next(reader, None)
        if not header:
            return []
        hanzi_index = _find_column(header, HANZI_HEADERS)
        note_index = _find_column(header, NOTE_HEADERS)
        tone_index = _find_column(header, TONE_HEADERS)
        if not _is_character_table(header, hanzi_index):
            raise ValueError(f"{source_path.name} 不是 #漢字 字表")
        if note_index is None:
            raise ValueError(f"{source_path.name} 找不到解釋/注釋列")

        for row_number, row in enumerate(reader, start=2):
            note = _cell(row, note_index).strip()
            if not note:
                continue
            records.append(
                {
                    "id": f"{source_path.name}:{row_number}",
                    "原始汉字": _cell(row, hanzi_index).strip(),
                    "原始注释": note,
                    "来源文件": source_path.name,
                    "原始行号": row_number,
                    "_原字音标": _cell(row, tone_index).strip(),
                }
            )
    return records


def _public_input(record: dict[str, Any]) -> dict[str, Any]:
    return {
        key: record[key]
        for key in ("id", "原始汉字", "原始注释", "来源文件", "原始行号")
    }


def _input_fingerprint(record: dict[str, Any]) -> str:
    payload = {
        "prompt_version": PROMPT_VERSION,
        **{
            key: record[key]
            for key in ("id", "原始汉字", "原始注释", "来源文件", "原始行号")
        },
    }
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _write_text_atomic(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Optional[Path] = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
            temporary_file.write(text)
            temporary_file.flush()
            os.fsync(temporary_file.fileno())
        os.replace(temporary_path, path)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def _write_jsonl(path: Path, records: list[dict[str, Any]]) -> None:
    _write_text_atomic(
        path,
        "".join(
            json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n"
            for record in records
        ),
    )


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    records = []
    with path.open("r", encoding="utf-8") as input_file:
        for line in input_file:
            if not line.strip():
                continue
            records.append(json.loads(line))
    return records


def _string_list(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value.strip()] if value.strip() else []
    if not isinstance(value, list):
        raise ValueError("标准词必须是字符串或数组")
    return [str(item).strip() for item in value if str(item).strip()]


def _parse_json_content(content: str) -> Any:
    content = content.strip()
    if content.startswith("```"):
        content = re.sub(r"^```(?:json)?\s*", "", content)
        content = re.sub(r"\s*```$", "", content)
    return json.loads(content)


def _validate_response(
    response: Any,
    batch: list[dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    if not isinstance(response, dict):
        raise ValueError("LLM 返回值必须是 JSON 对象")
    results = response.get("results")
    if not isinstance(results, list):
        raise ValueError("LLM 返回值缺少 results 数组")

    expected_ids = {item["id"] for item in batch}
    parsed = {}
    for item in results:
        if not isinstance(item, dict):
            raise ValueError("results 中包含非对象")
        record_id = str(item.get("id", "")).strip()
        if record_id not in expected_ids:
            raise ValueError(f"LLM 返回未知 id: {record_id}")
        if record_id in parsed:
            raise ValueError(f"LLM 重复返回 id: {record_id}")
        entries = item.get("entries")
        if not isinstance(entries, list):
            raise ValueError(f"{record_id} 的 entries 不是数组")

        normalized_entries = []
        for index, entry in enumerate(entries):
            if not isinstance(entry, dict):
                raise ValueError(f"{record_id} entries[{index}] 不是对象")
            dialect_word = str(entry.get("方言词", "")).strip()
            if not dialect_word:
                raise ValueError(f"{record_id} entries[{index}] 缺少方言词")
            confidence = entry.get("置信度")
            if not isinstance(confidence, (int, float)) or not 0 <= confidence <= 1:
                raise ValueError(f"{record_id} entries[{index}] 置信度无效")
            normalized_entries.append(
                {
                    "方言词": dialect_word,
                    "标准词": _string_list(entry.get("标准词", [])),
                    "置信度": confidence,
                }
            )
        parsed[record_id] = {"id": record_id, "entries": normalized_entries}

    missing_ids = expected_ids - parsed.keys()
    if missing_ids:
        raise ValueError(f"LLM 缺少 id: {','.join(sorted(missing_ids))}")
    return parsed


def _call_deepseek(
    api_key: str,
    model: str,
    batch: list[dict[str, Any]],
    *,
    retries: int = 4,
) -> dict[str, dict[str, Any]]:
    user_content = (
        "输入记录如下。每条记录必须独立处理：\n"
        + json.dumps(batch, ensure_ascii=False, separators=(",", ":"))
    )
    body = json.dumps(
        {
            "model": model,
            "messages": [
                {"role": "system", "content": PROMPT},
                {"role": "user", "content": user_content},
            ],
            "response_format": {"type": "json_object"},
            "temperature": 0.1,
        },
        ensure_ascii=False,
    ).encode("utf-8")

    last_error: Optional[Exception] = None
    for attempt in range(retries):
        request = urllib.request.Request(
            API_URL,
            data=body,
            headers={
                "Content-Type": "application/json",
                "Authorization": f"Bearer {api_key}",
            },
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=180) as response:
                payload = json.loads(response.read().decode("utf-8"))
            content = payload["choices"][0]["message"]["content"]
            return _validate_response(_parse_json_content(content), batch)
        except urllib.error.HTTPError as error:
            detail = error.read().decode("utf-8", errors="replace")
            message = f"DeepSeek HTTP {error.code}: {detail[:800]}"
            if error.code in {400, 401, 402, 403}:
                raise RuntimeError(message) from error
            last_error = RuntimeError(message)
        except (
            urllib.error.URLError,
            TimeoutError,
            OSError,
            KeyError,
            ValueError,
            json.JSONDecodeError,
        ) as error:
            last_error = error
        if attempt < retries - 1:
            time.sleep(min(2**attempt + random.random(), 20))
    raise RuntimeError(f"DeepSeek 重试耗尽: {last_error!r}")


def _load_done(result_path: Path) -> dict[str, dict[str, Any]]:
    return {
        record["id"]: record
        for record in _read_jsonl(result_path)
        if isinstance(record, dict) and record.get("id")
    }


def _attach_input_fingerprints(
    response: dict[str, dict[str, Any]],
    batch: list[dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    records_by_id = {record["id"]: record for record in batch}
    return {
        record_id: {
            **result,
            "input_fingerprint": _input_fingerprint(records_by_id[record_id]),
        }
        for record_id, result in response.items()
    }


def _ordered_results(
    records: list[dict[str, Any]],
    results: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    return [
        results[record["id"]]
        for record in records
        if record["id"] in results
    ]


def _write_workbook(
    point: str,
    records: list[dict[str, Any]],
    results: dict[str, dict[str, Any]],
    errors: dict[str, str],
    xlsx_path: Path,
) -> None:
    workbook = Workbook()
    word_sheet = workbook.active
    word_sheet.title = "词表"
    word_sheet.append(WORD_SHEET_HEADERS)

    for record in records:
        result = results.get(record["id"], {"entries": []})
        for entry in result.get("entries", []):
            standards = entry.get("标准词", []) or [""]
            for standard_word in standards:
                word_sheet.append(
                    [
                        standard_word,
                        entry["方言词"],
                        "",
                        record["原始注释"],
                        "",
                        entry["置信度"],
                        record["原始汉字"],
                        record["_原字音标"],
                        record["原始注释"],
                        record["来源文件"],
                        record["原始行号"],
                    ]
                )

    row_sheet = workbook.create_sheet("逐行结果")
    row_sheet.append(ROW_SHEET_HEADERS)
    for record in records:
        record_id = record["id"]
        if record_id in errors:
            status = "失败"
            result_text = json.dumps(
                {"error": errors[record_id]},
                ensure_ascii=False,
                separators=(",", ":"),
            )
        elif record_id not in results:
            status = "待处理"
            result_text = ""
        else:
            entries = results[record_id].get("entries", [])
            status = "已生成" if entries else "无词条"
            result_text = json.dumps(
                entries,
                ensure_ascii=False,
                separators=(",", ":"),
            )
        row_sheet.append(
            [
                record_id,
                record["原始汉字"],
                record["原始注释"],
                record["来源文件"],
                record["原始行号"],
                status,
                result_text,
            ]
        )

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            column_index = column_cells[0].column
            max_length = max(
                len(str(cell.value or "")) for cell in column_cells[:200]
            )
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_length + 2, 10),
                45,
            )

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)


def run_analysis(
    source_dir: Path = DEFAULT_SOURCE_DIR,
    *,
    points: Optional[list[str]] = None,
    batch_size: int = 20,
    workers: int = 4,
    output_dir: Path = DEFAULT_OUTPUT_DIR,
    request_fn: Optional[
        Callable[[str, list[dict[str, Any]]], dict[str, dict[str, Any]]]
    ] = None,
    api_key: Optional[str] = None,
    model: str = "deepseek-chat",
    reset: bool = False,
) -> dict[str, dict[str, int | str]]:
    if batch_size <= 0 or workers <= 0:
        raise ValueError("batch_size 和 workers 必须大于 0")
    if request_fn is None:
        api_key = _load_api_key(api_key)
        if not api_key:
            raise RuntimeError(
                "请在环境变量 DEEPSEEK_API_KEY 或项目 .env 中设置 DeepSeek API key"
            )

    source_dir = Path(source_dir)
    output_dir = Path(output_dir)
    input_dir = output_dir / "input"
    result_dir = output_dir / "result"
    xlsx_dir = output_dir / "xlsx"
    if reset:
        _clear_dir(output_dir)
    else:
        input_dir.mkdir(parents=True, exist_ok=True)
        result_dir.mkdir(parents=True, exist_ok=True)
        xlsx_dir.mkdir(parents=True, exist_ok=True)

    selected_points = points or list(DEFAULT_POINTS)
    summary = {}

    for point in selected_points:
        records = _source_records(source_dir, point)
        public_records = [_public_input(record) for record in records]
        input_path = input_dir / f"{point}.jsonl"
        _write_jsonl(input_path, public_records)

        result_path = result_dir / f"{point}.jsonl"
        cached = _load_done(result_path)
        current_fingerprints = {
            record["id"]: _input_fingerprint(record)
            for record in public_records
        }
        done = {
            record_id: result
            for record_id, result in cached.items()
            if result.get("input_fingerprint")
            == current_fingerprints.get(record_id)
        }
        errors: dict[str, str] = {}
        pending = [
            record
            for record in public_records
            if record["id"] not in done
        ]
        batches = [
            pending[index:index + batch_size]
            for index in range(0, len(pending), batch_size)
        ]

        def request_batch(batch: list[dict[str, Any]]):
            if request_fn is not None:
                return request_fn(point, batch)
            return _call_deepseek(api_key or "", model, batch)

        lock = threading.Lock()

        def process_batch(batch: list[dict[str, Any]]):
            try:
                response = request_batch(batch)
                if request_fn is not None:
                    response = _validate_response(response, batch)
                return batch, response, None
            except Exception as error:
                return batch, None, str(error)

        with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as executor:
            futures = [executor.submit(process_batch, batch) for batch in batches]
            for future in concurrent.futures.as_completed(futures):
                batch, response, error = future.result()
                if error:
                    if any(
                        marker in error
                        for marker in (
                            "DeepSeek HTTP 400",
                            "DeepSeek HTTP 401",
                            "DeepSeek HTTP 402",
                            "DeepSeek HTTP 403",
                        )
                    ):
                        raise RuntimeError(error)
                    for record in batch:
                        errors[record["id"]] = error
                    continue
                with lock:
                    done.update(_attach_input_fingerprints(response, batch))
                    _write_jsonl(result_path, _ordered_results(public_records, done))

        all_results = {
            record["id"]: done[record["id"]]
            for record in public_records
            if record["id"] in done
        }
        _write_jsonl(result_path, _ordered_results(public_records, all_results))
        error_path = result_dir / f"{point}.errors.jsonl"
        if errors:
            _write_jsonl(
                error_path,
                [
                    {"id": record_id, "error": message}
                    for record_id, message in errors.items()
                ],
            )
        elif error_path.exists():
            error_path.unlink()

        result_by_id = {
            record["id"]: all_results[record["id"]]
            for record in public_records
            if record["id"] in all_results
        }
        _write_workbook(
            point,
            records,
            result_by_id,
            errors,
            xlsx_dir / f"{point}.xlsx",
        )
        summary[point] = {
            "status": (
                "partial"
                if errors or len(all_results) < len(public_records)
                else "ok"
            ),
            "rows": len(records),
            "batches": len(batches),
            "completed": len(all_results),
            "errors": len(errors),
            "word_rows": sum(
                len(result.get("entries", [])) for result in result_by_id.values()
            ),
        }

    _write_text_atomic(
        output_dir / "manifest.json",
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
    )
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(
        description="将音典字表逐行交给 DeepSeek 分析，并按地点生成 XLSX 词表"
    )
    parser.add_argument("command", choices=["run"])
    parser.add_argument(
        "--source-dir",
        type=Path,
        default=DEFAULT_SOURCE_DIR,
        help="原始音典 TSV 目录",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="输入快照、JSONL 结果和 XLSX 输出目录",
    )
    parser.add_argument(
        "--points",
        default=",".join(DEFAULT_POINTS),
        help="地点，逗号分隔",
    )
    parser.add_argument("--batch-size", type=int, default=20)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--model", default="deepseek-chat")
    parser.add_argument("--reset", action="store_true", help="清空本次输出后重新运行")
    args = parser.parse_args()

    summary = run_analysis(
        args.source_dir,
        points=[point.strip() for point in args.points.split(",") if point.strip()],
        batch_size=args.batch_size,
        workers=args.workers,
        output_dir=args.output_dir,
        model=args.model,
        reset=args.reset,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    if any(item["status"] != "ok" for item in summary.values()):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
